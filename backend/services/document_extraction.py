from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument

from backend.core.exceptions import DocumentExtractionError


DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


@dataclass(frozen=True)
class ExtractedPage:
    page_number: int
    text: str


@dataclass(frozen=True)
class ExtractedDocument:
    pages: tuple[ExtractedPage, ...]

    @property
    def page_count(self) -> int:
        return len(self.pages)

    @property
    def text(self) -> str:
        return "\n".join(
            page.text
            for page in self.pages
        ).strip()

    @property
    def character_count(self) -> int:
        return len(self.text)

    @property
    def word_count(self) -> int:
        return len(self.text.split())


def extract_pdf_document(
    content: bytes,
) -> ExtractedDocument:
    try:
        with pymupdf.open(
            stream=content,
            filetype="pdf",
        ) as pdf:
            pages = tuple(
                ExtractedPage(
                    page_number=page_number,
                    text=page.get_text(),
                )
                for page_number, page in enumerate(
                    pdf,
                    start=1,
                )
            )

        return ExtractedDocument(
            pages=pages,
        )

    except Exception as exc:
        raise DocumentExtractionError(
            "Failed to extract PDF text"
        ) from exc


def extract_docx_document(
    content: bytes,
) -> ExtractedDocument:
    try:
        document = DocxDocument(
            BytesIO(content)
        )

        parts = []

        paragraphs_by_element = {
            paragraph._p: paragraph
            for paragraph in document.paragraphs
        }

        tables_by_element = {
            table._tbl: table
            for table in document.tables
        }

        for element in document.element.body:
            if element in paragraphs_by_element:
                paragraph = paragraphs_by_element[element]

                if paragraph.text.strip():
                    parts.append(
                        paragraph.text.strip()
                    )

            elif element in tables_by_element:
                table = tables_by_element[element]

                for row in table.rows:
                    cells = [
                        cell.text.strip()
                        for cell in row.cells
                    ]

                    row_text = " | ".join(
                        cell
                        for cell in cells
                        if cell
                    )

                    if row_text:
                        parts.append(row_text)

        # DOCX has no fixed pagination, so the body is one logical page.
        page = ExtractedPage(
            page_number=1,
            text="\n".join(parts),
        )

        return ExtractedDocument(
            pages=(page,),
        )

    except Exception as exc:
        raise DocumentExtractionError(
            "Failed to extract DOCX text"
        ) from exc


def extract_document(
    filename: str,
    content_type: str | None,
    content: bytes,
) -> ExtractedDocument:
    extension = Path(filename).suffix.lower()

    if (
        extension == ".pdf"
        or content_type == "application/pdf"
    ):
        return extract_pdf_document(content)

    if (
        extension == ".docx"
        or content_type == DOCX_CONTENT_TYPE
    ):
        return extract_docx_document(content)

    raise DocumentExtractionError(
        "Unsupported document format"
    )
