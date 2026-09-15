from pathlib import Path

from docx import Document


OUTPUT_DIR = Path("benchmark-data")


DOCUMENTS = {
    "leave_policy.docx": (
        "Leave Policy",
        [
            (
                "Annual Leave",
                "Employees receive 20 days of annual leave per year.",
            ),
            (
                "Sick Leave",
                "Employees receive 10 days of sick leave per year.",
            ),
            (
                "Holiday Leave",
                "The company observes 12 paid public holidays each year.",
            ),
            (
                "Leave Approval",
                "Leave requests longer than three working days require manager approval.",
            ),
        ],
    ),
    "remote_work_policy.docx": (
        "Remote Work Policy",
        [
            (
                "Eligibility",
                "Employees who have completed probation may request remote work.",
            ),
            (
                "Weekly Limit",
                "Eligible employees may work remotely up to three days per week.",
            ),
            (
                "Core Hours",
                "Remote employees must be available from 10 AM to 4 PM.",
            ),
            (
                "Equipment and Security",
                "The company provides a laptop and required security software for approved remote employees.",
            ),
        ],
    ),
    "information_security_policy.docx": (
        "Information Security Policy",
        [
            (
                "Multi-Factor Authentication",
                "All employees must use multi-factor authentication for supported company systems.",
            ),
            (
                "Data Classification",
                "Confidential information must not be shared outside the company without authorization.",
            ),
            (
                "Device Security",
                "Company devices must use disk encryption and approved security software.",
            ),
            (
                "Incident Reporting",
                "Security incidents must be reported immediately to the security team.",
            ),
        ],
    ),
    "password_authentication_policy.docx": (
        "Password Authentication Policy",
        [
            (
                "Password Length",
                "Passwords must contain at least 12 characters.",
            ),
            (
                "Password Rotation",
                "Passwords must be changed every 90 days.",
            ),
            (
                "Password Reuse",
                "Employees may not reuse their previous five passwords.",
            ),
            (
                "Account Lockout and Recovery",
                "Accounts are locked after five failed login attempts and password recovery requires identity verification.",
            ),
        ],
    ),
    "expense_policy.docx": (
        "Expense Policy",
        [
            (
                "Meal Expense",
                "Employees may claim up to Rs. 800 per day for eligible meals incurred during approved business activity.",
            ),
            (
                "Receipt Requirement",
                "Expenses above Rs. 500 require an itemized receipt.",
            ),
            (
                "Submission Deadline",
                "Expense reports must be submitted within 30 days.",
            ),
            (
                "Approval and Currency",
                "All reimbursement claims require manager approval.",
            ),
        ],
    ),
    "attendance_policy.docx": (
        "Attendance Policy",
        [
            (
                "Work Week",
                "The standard work week consists of five working days.",
            ),
            (
                "Late Arrival",
                "Employees arriving more than 30 minutes late must inform their manager.",
            ),
            (
                "Attendance Recording",
                "Employees must record attendance through the company system.",
            ),
            (
                "Absence and Flexible Hours",
                "Unplanned absence must be reported before the start of the workday.",
            ),
        ],
    ),
    "travel_policy.docx": (
        "Travel Policy",
        [
            (
                "Flight Class",
                "Domestic business travel is normally booked in economy class.",
            ),
            (
                "Hotel Limit",
                "Employees may claim up to Rs. 5000 per night for eligible hotels.",
            ),
            (
                "Advance Booking",
                "Travel should be booked at least seven days in advance.",
            ),
            (
                "Ground Transport and Approval",
                "Reasonable taxi and local transport costs are reimbursable for approved business travel.",
            ),
        ],
    ),
    "employee_conduct_policy.docx": (
        "Employee Conduct Policy",
        [
            (
                "Professional Conduct",
                "Employees must communicate respectfully with colleagues and professional contacts.",
            ),
            (
                "Harassment and Discrimination",
                "Harassment and discrimination are prohibited.",
            ),
            (
                "Conflicts of Interest",
                "Employees must disclose potential conflicts of interest.",
            ),
            (
                "Confidentiality and Discipline",
                "Employees must protect confidential company information.",
            ),
        ],
    ),
    "workplace_safety_policy.docx": (
        "Workplace Safety Policy",
        [
            (
                "Emergency Exits",
                "Emergency exits must remain clear and accessible at all times.",
            ),
            (
                "Fire Safety",
                "Employees must follow building fire evacuation procedures.",
            ),
            (
                "Protective Equipment",
                "Required protective equipment must be worn in designated areas.",
            ),
            (
                "Hazard Reporting and First Aid",
                "Workplace hazards must be reported to facilities immediately.",
            ),
        ],
    ),
    "emergency_alert_system.docx": (
        "Emergency Alert System",
        [
            (
                "Microcontroller",
                "The emergency alert system uses an ESP32 microcontroller as the primary control unit.",
            ),
            (
                "Gas and Flame Detection",
                "An MQ-2 gas sensor and an IR flame sensor are used for gas and fire detection.",
            ),
            (
                "Wireless Communication",
                "LoRa modules provide wireless communication between the transmitter and receiver.",
            ),
            (
                "Escalation",
                "If an alert is not acknowledged within two minutes, the system escalates the notification to the next contact.",
            ),
        ],
    ),
}


EXPANSION_TEMPLATES = [
    "This section defines the standard expectations that employees should follow in normal business situations. The policy is intended to provide a consistent process and reduce uncertainty when employees need to make decisions.",
    "Employees should follow the approved company procedure and use the designated systems whenever the policy requires a formal request, notification, approval, or record. Managers are responsible for helping employees understand the process.",
    "Exceptions may occur because of operational requirements, urgent circumstances, or special business situations. Where an exception is necessary, the employee should communicate the situation promptly and obtain the appropriate approval.",
    "Records associated with this policy should be accurate and complete. Employees should provide enough information for a reviewer to understand the action taken, the business reason, and any relevant supporting details.",
    "Managers should apply the policy consistently while considering reasonable operational circumstances. Questions about interpretation should be raised through the normal management or administrative channel.",
    "The policy should be reviewed periodically as business requirements change. Employees are responsible for following the current version communicated by the company rather than relying on an outdated copy.",
]


def create_document(filename, title, sections):
    document = Document()

    document.add_heading(
        title,
        level=1,
    )

    for heading, key_fact in sections:
        document.add_heading(
            heading,
            level=2,
        )

        # Put the key fact first so the important retrieval signal
        # remains easy to identify.
        document.add_paragraph(key_fact)

        for template in EXPANSION_TEMPLATES:
            document.add_paragraph(
                f"{template} {key_fact}"
            )

        # Repeat the core policy fact in a final explanatory paragraph.
        document.add_paragraph(
            f"Important policy requirement: {key_fact}"
        )

    document.save(
        OUTPUT_DIR / filename
    )


def main():
    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    for filename, (title, sections) in DOCUMENTS.items():
        create_document(
            filename,
            title,
            sections,
        )

    print(
        f"Created {len(DOCUMENTS)} benchmark documents."
    )


if __name__ == "__main__":
    main()
