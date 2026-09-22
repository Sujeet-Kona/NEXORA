from io import BytesIO

import fitz
from docx import Document as DocxDocument
import pytest

from backend.core.exceptions import DocumentExtractionError
from backend.services.document_extraction import (
    extract_docx_document,
    extract_document,
    extract_pdf_document,
)


def make_pdf() -> bytes:
    document = fitz.open()

    page = document.new_page()

    page.insert_text(
        (72, 72),
        "Nexora HR Policy",
    )

    page.insert_text(
        (72, 100),
        "Employees receive 20 days of annual leave.",
    )

    content = document.tobytes()

    document.close()

    return content


def make_multi_page_pdf() -> bytes:
    document = fitz.open()

    first_page = document.new_page()

    first_page.insert_text(
        (72, 72),
        "First page content",
    )

    second_page = document.new_page()

    second_page.insert_text(
        (72, 72),
        "Second page content",
    )

    content = document.tobytes()

    document.close()

    return content


def make_image_only_pdf() -> bytes:
    document = fitz.open()

    page = document.new_page(
        width=200,
        height=200,
    )

    pixmap = fitz.Pixmap(
        fitz.csRGB,
        fitz.IRect(0, 0, 50, 50),
        False,
    )

    page.insert_image(
        fitz.Rect(0, 0, 200, 200),
        pixmap=pixmap,
    )

    content = document.tobytes()

    document.close()

    return content


def make_docx() -> bytes:
    document = DocxDocument()

    document.add_paragraph(
        "Nexora Security Policy"
    )

    document.add_paragraph(
        "All employees must use MFA."
    )

    table = document.add_table(
        rows=2,
        cols=2,
    )

    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Purpose"
    table.cell(1, 0).text = "MFA"
    table.cell(1, 1).text = "Account protection"

    buffer = BytesIO()

    document.save(buffer)

    return buffer.getvalue()


DOCX_TEXT = (
    "Nexora Security Policy\n"
    "All employees must use MFA.\n"
    "Component | Purpose\n"
    "MFA | Account protection"
)


def test_extract_pdf_document_returns_numbered_pages():
    extracted = extract_pdf_document(
        make_pdf(),
    )

    assert extracted.page_count == 1
    assert len(extracted.pages) == 1
    assert extracted.pages[0].page_number == 1
    assert "Nexora HR Policy" in extracted.pages[0].text
    assert (
        "Employees receive 20 days of annual leave."
        in extracted.pages[0].text
    )


def test_extract_pdf_document_separates_page_text():
    extracted = extract_pdf_document(
        make_multi_page_pdf(),
    )

    assert extracted.page_count == 2

    assert [
        page.page_number
        for page in extracted.pages
    ] == [1, 2]

    assert "First page content" in extracted.pages[0].text
    assert "Second page content" not in extracted.pages[0].text

    assert "Second page content" in extracted.pages[1].text
    assert "First page content" not in extracted.pages[1].text


def test_flat_text_is_derived_from_pages():
    extracted = extract_pdf_document(
        make_multi_page_pdf(),
    )

    assert extracted.text == (
        "\n".join(
            page.text
            for page in extracted.pages
        ).strip()
    )

    assert not extracted.text.startswith("\n")
    assert not extracted.text.endswith("\n")


def test_pdf_stats_match_flat_text():
    extracted = extract_pdf_document(
        make_pdf(),
    )

    assert extracted.word_count == 10
    assert extracted.character_count == len(extracted.text)
    assert extracted.character_count == len(
        "Nexora HR Policy\n"
        "Employees receive 20 days of annual leave."
    )


def test_image_only_pdf_reports_zero_text_stats():
    extracted = extract_pdf_document(
        make_image_only_pdf(),
    )

    assert extracted.page_count == 1
    assert extracted.text == ""
    assert extracted.word_count == 0
    assert extracted.character_count == 0


def test_extract_docx_document_is_one_logical_page():
    extracted = extract_docx_document(
        make_docx(),
    )

    assert extracted.page_count == 1
    assert extracted.pages[0].page_number == 1
    assert extracted.pages[0].text == DOCX_TEXT
    assert extracted.text == DOCX_TEXT
    assert extracted.word_count == 15
    assert extracted.character_count == len(DOCX_TEXT)


def test_extract_docx_document_includes_tables():
    extracted = extract_docx_document(
        make_docx(),
    )

    assert "Component" in extracted.text
    assert "Purpose" in extracted.text
    assert "MFA" in extracted.text
    assert "Account protection" in extracted.text


def test_extract_document_detects_pdf():
    extracted = extract_document(
        filename="policy.pdf",
        content_type="application/pdf",
        content=make_pdf(),
    )

    assert extracted.page_count == 1
    assert "Nexora HR Policy" in extracted.text


def test_extract_document_detects_docx():
    extracted = extract_document(
        filename="policy.docx",
        content_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        content=make_docx(),
    )

    assert extracted.page_count == 1
    assert "Nexora Security Policy" in extracted.text


def test_invalid_pdf_raises():
    with pytest.raises(DocumentExtractionError):
        extract_pdf_document(
            b"not a real pdf"
        )


def test_unsupported_format_raises():
    with pytest.raises(DocumentExtractionError):
        extract_document(
            filename="data.exe",
            content_type="application/octet-stream",
            content=b"bad",
        )
